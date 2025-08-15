import openpyxl
from django.contrib.auth import get_user_model
from django.db import transaction

from .models import Action, Plan, ActionPlan
from .models_import import SyncReport

User = get_user_model()


def _get_cell_value(row, col_letter):
    index = ord(col_letter.upper()) - 65
    if index < 0 or index >= len(row):
        return None
    return row[index].value


@transaction.atomic
def import_plan_from_excel(plan: Plan, file_path: str, profile):
    wb = openpyxl.load_workbook(file_path)
    if profile.sheet_name not in wb.sheetnames:
        raise ValueError(f"Feuille {profile.sheet_name} introuvable.")
    ws = wb[profile.sheet_name]

    created = updated = skipped = 0
    errors = []

    for row in ws.iter_rows(min_row=profile.header_row + 1):
        try:
            row_data = {}
            for col_letter, field_name in profile.mapping.items():
                row_data[field_name] = _get_cell_value(row, col_letter)

            code = str(row_data.get("code") or "").strip()
            if not code:
                skipped += 1
                continue

            action, created_flag = Action.objects.get_or_create(
                code=code,
                defaults={
                    "title": row_data.get("title", ""),
                    "description": row_data.get("description", ""),
                    "priority": row_data.get("priority", "MED"),
                    "status": row_data.get("status", "A_FAIRE"),
                    "due_date": row_data.get("due_date"),
                    "created_by": plan.owner,
                },
            )

            if created_flag:
                created += 1
            else:
                action.title = row_data.get("title", action.title)
                action.description = row_data.get("description", action.description)
                action.priority = row_data.get("priority", action.priority)
                action.status = row_data.get("status", action.status)
                action.due_date = row_data.get("due_date", action.due_date)
                action.save()
                updated += 1

            ActionPlan.objects.get_or_create(action=action, plan=plan)

            responsables_str = row_data.get("responsables")
            if responsables_str:
                for resp_name in responsables_str.split(";"):
                    resp_name = resp_name.strip()
                    if not resp_name:
                        continue
                    if profile.allow_user_creation:
                        user, _ = User.objects.get_or_create(username=resp_name)
                    else:
                        user = User.objects.filter(username=resp_name).first()
                    if user:
                        action.responsables.add(user)
        except Exception as exc:  # pylint: disable=broad-except
            errors.append({"row": row[0].row, "error": str(exc)})

    report = SyncReport.objects.create(
        plan=plan,
        created_count=created,
        updated_count=updated,
        skipped_count=skipped,
        errors=errors,
    )
    return report
